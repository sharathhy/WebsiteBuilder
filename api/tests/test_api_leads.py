import csv
import io
import json
from pathlib import Path

import httpx
from openpyxl import load_workbook

from conftest import ADMIN, REVIEWER, login
from leadforge.models import User
from leadforge.pipeline.runner import InlineDispatcher

FILLED_AGENCY = """\
agency_name: Test Studio
sender_name: Test Sender
sender_email: hello@studio.example
physical_address: 1 Test Street, Austin, TX 78701, USA
booking_url: https://cal.example/test
monthly_budget_usd: 50
pricing:
  IN: {currency: INR, symbol: "₹", tiers: {starter: 15000, growth: 35000, premium: 75000}}
  US: {currency: USD, symbol: "$", tiers: {starter: 499, growth: 1299, premium: 2999}}
"""


async def _run(client: httpx.AsyncClient, inline: InlineDispatcher, h: dict[str, str]) -> str:
    resp = await client.post(
        "/api/campaigns",
        headers=h,
        json={"name": "Austin dentists", "cities": ["austin"], "categories": ["dentists"],
              "max_results_per_target": 6},
    )  # fmt: skip
    assert resp.status_code == 201, resp.text
    cid = resp.json()["id"]
    run = await client.post(f"/api/campaigns/{cid}/run", headers=h)
    assert run.status_code == 202 and len(run.json()) == 1
    await inline.drain()
    return str(cid)


async def test_campaign_rejects_unknown_targets(
    client: httpx.AsyncClient, users: dict[str, User]
) -> None:
    h = await login(client, ADMIN)
    resp = await client.post(
        "/api/campaigns", headers=h,
        json={"name": "x", "cities": ["atlantis"], "categories": ["dentists"]},
    )  # fmt: skip
    assert resp.status_code == 422
    rh = await login(client, REVIEWER)
    resp = await client.post(
        "/api/campaigns", headers=rh,
        json={"name": "x", "cities": ["austin"], "categories": ["dentists"]},
    )  # fmt: skip
    assert resp.status_code == 403


async def test_leads_list_detail_and_export(
    client: httpx.AsyncClient, users: dict[str, User], inline: InlineDispatcher
) -> None:
    h = await login(client, ADMIN)
    cid = await _run(client, inline, h)

    progress = (await client.get(f"/api/campaigns/{cid}/progress", headers=h)).json()
    assert progress["total_leads"] == 6
    assert progress["jobs"]["discover"] == {"succeeded": 1}

    page = (await client.get("/api/leads", headers=h, params={"campaign_id": cid})).json()
    assert page["total"] == 6
    scores = [r["score"] for r in page["items"] if r["score"] is not None]
    assert scores == sorted(scores)  # default sort: worst first
    for r in page["items"]:
        assert r["email"] in ("not found",) or "@" in r["email"]
        assert r["phone"].startswith("+1")

    low = (await client.get("/api/leads", headers=h, params={"max_score": 59})).json()
    assert all(r["score"] <= 59 for r in low["items"])
    assert all(r["preview_url"] for r in low["items"])
    no_site = (await client.get("/api/leads", headers=h, params={"has_website": False})).json()
    assert all(r["website"] == "not found" for r in no_site["items"])

    lead_id = low["items"][0]["id"]
    detail = (await client.get(f"/api/leads/{lead_id}", headers=h)).json()
    assert detail["sites"][0]["status"] == "live"
    assert detail["sites"][0]["score"] >= 90
    assert detail["drafts"]
    assert any(a["generated_site_id"] for a in detail["audits"])  # the re-score of the preview

    resp = await client.get("/api/leads/export", headers=h, params={"campaign_id": cid})
    assert resp.status_code == 200 and resp.headers["content-type"].startswith("text/csv")
    rows = list(csv.DictReader(io.StringIO(resp.text)))
    assert len(rows) == 6 and {"email", "email_source_url", "score", "preview_url"} <= set(rows[0])

    resp = await client.get("/api/leads/export", headers=h, params={"format": "xlsx"})
    ws = load_workbook(io.BytesIO(resp.content)).active
    assert ws is not None and ws.max_row == 7

    # Previews are served locally (mock deploy target) with noindex headers.
    preview = detail["sites"][0]["preview_url"].removeprefix("http://test")
    resp = await client.get(preview)
    assert resp.status_code == 200 and resp.headers["x-robots-tag"] == "noindex, nofollow"


async def test_approval_flow_sends_only_after_a_human_approves(
    client: httpx.AsyncClient,
    users: dict[str, User],
    inline: InlineDispatcher,
    config_dir: Path,
    tmp_path: Path,
) -> None:
    h = await login(client, REVIEWER)
    await _run(client, inline, await login(client, ADMIN))
    outbox = tmp_path / "storage" / "outbox"

    drafts = (await client.get("/api/drafts", headers=h, params={"channel": "email"})).json()
    assert drafts and all(d["status"] == "drafted" for d in drafts)
    with_email = [d for d in drafts if d["to"]]
    assert with_email, "fixtures should include businesses with a published email"
    draft = with_email[0]

    # The cron sender ignores unapproved drafts.
    from leadforge.models.enums import JobQueue
    from leadforge.services.jobs import create_job

    async with inline.sessionmaker() as db:
        job, _ = await create_job(
            db, queue=JobQueue.EMAIL, kind="send_due", idempotency_key="cron1"
        )
    await inline.dispatch(job)
    await inline.drain()
    assert not outbox.exists()

    # Placeholder agency details block sending.
    resp = await client.post(
        f"/api/drafts/{draft['id']}/approve", headers=h, json={"send_now": True}
    )
    assert resp.status_code == 409 and "agency.yaml" in resp.json()["detail"]

    # Edits keep the unsubscribe link even if the reviewer deletes it.
    edited = await client.patch(
        f"/api/drafts/{draft['id']}", headers=h, json={"body": "Hi there, short note."}
    )
    assert edited.status_code == 200 and "/u/" in edited.json()["body"]

    (config_dir / "agency.yaml").write_text(FILLED_AGENCY)
    resp = await client.post(
        f"/api/drafts/{draft['id']}/approve", headers=h, json={"send_now": True}
    )
    assert resp.status_code == 200 and resp.json()["status"] == "approved"
    await inline.drain()
    sent = [json.loads(p.read_text()) for p in outbox.glob("*.json")]
    assert len(sent) == 1
    assert sent[0]["headers"]["List-Unsubscribe-Post"] == "List-Unsubscribe=One-Click"
    assert "1 Test Street" in sent[0]["text"]

    detail = (await client.get(f"/api/leads/{draft['business_id']}", headers=h)).json()
    statuses = {(d["sequence_day"], d["variant"]): d["status"] for d in detail["drafts"]}
    assert statuses[(0, draft["variant"])] == "sent"
    assert "cancelled" in statuses.values()  # the other A/B variant
    followups = [d for d in detail["drafts"] if d["sequence_day"] in (3, 7)]
    assert len(followups) == 2 and all(d["status"] == "drafted" for d in followups)

    # One-click unsubscribe: suppresses the address and cancels the follow-ups.
    token_body = sent[0]["headers"]["List-Unsubscribe"].strip("<>").removeprefix("http://test")
    page = await client.get(token_body)
    assert page.status_code == 200 and "<form" in page.text
    resp = await client.post(token_body)
    assert resp.status_code == 200
    detail = (await client.get(f"/api/leads/{draft['business_id']}", headers=h)).json()
    assert detail["lead"]["lead_status"] == "suppressed"
    assert all(d["status"] != "drafted" for d in detail["drafts"])
    supp = (await client.get("/api/suppression", headers=h)).json()
    assert any(e["reason"] == "unsubscribe" for e in supp)


async def test_manual_suppression_blocks_approval(
    client: httpx.AsyncClient, users: dict[str, User], inline: InlineDispatcher
) -> None:
    h = await login(client, ADMIN)
    await _run(client, inline, h)
    drafts = (await client.get("/api/drafts", headers=h, params={"channel": "email"})).json()
    draft = next(d for d in drafts if d["to"])
    resp = await client.post("/api/suppression", headers=h, json={"email": draft["to"]})
    assert resp.status_code == 201
    resp = await client.post(
        f"/api/drafts/{draft['id']}/approve", headers=h, json={"send_now": False}
    )
    assert resp.status_code == 409


async def test_contact_form_and_honeypot(
    client: httpx.AsyncClient, users: dict[str, User], inline: InlineDispatcher
) -> None:
    h = await login(client, ADMIN)
    await _run(client, inline, h)
    low = (await client.get("/api/leads", headers=h, params={"max_score": 59})).json()
    detail = (await client.get(f"/api/leads/{low['items'][0]['id']}", headers=h)).json()
    site_id = detail["sites"][0]["id"]

    resp = await client.post(
        f"/api/public/forms/{site_id}",
        json={"name": "Pat", "email": "pat@visitor.example", "message": "Do you open Saturdays?"},
    )
    assert resp.status_code == 202 and resp.headers["access-control-allow-origin"] == "*"
    resp = await client.post(f"/api/public/forms/{site_id}", data={"name": "Pat", "message": "hi"})
    assert resp.status_code == 422  # needs an email or phone
    resp = await client.post(
        f"/api/public/forms/{site_id}",
        data={"name": "Bot", "email": "b@b.example", "message": "x", "company_website": "spam"},
    )
    assert resp.status_code == 202

    from sqlalchemy import func, select

    from leadforge.models import FormSubmission

    async with inline.sessionmaker() as db:
        assert await db.scalar(select(func.count()).select_from(FormSubmission)) == 1


async def test_analytics(
    client: httpx.AsyncClient, users: dict[str, User], inline: InlineDispatcher
) -> None:
    h = await login(client, ADMIN)
    await _run(client, inline, h)
    overview = (await client.get("/api/analytics/overview", headers=h)).json()
    assert overview["funnel"][0] == {"stage": "Discovered", "count": 6}
    costs = (await client.get("/api/analytics/costs", headers=h)).json()
    assert costs["spent_usd"] == 0 and costs["mock_mode"] is True
    assert costs["providers"]["google_places"]["mocked_calls"] >= 1
    queues = (await client.get("/api/analytics/queues", headers=h)).json()
    assert queues["dead_letters"] == []
